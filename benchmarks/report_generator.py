"""Write timestamped JSON, CSV, Markdown, and Excel benchmark reports."""

from __future__ import annotations

import csv
import json
import platform
import subprocess
from datetime import datetime
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from typing import Any

from benchmarks import __version__ as benchmark_version
from benchmarks.config import BenchmarkConfig, BenchmarkResult
from benchmarks.resource_monitor import system_information
from benchmarks.runtime import model_inventory


_PACKAGES = (
    "fastapi", "pydantic", "langgraph", "spacy", "scispacy", "transformers",
    "torch", "sentence-transformers", "chromadb", "pymupdf", "pytesseract",
    "pillow", "psutil", "openpyxl",
)


def _git_value(*args: str) -> str | None:
    try:
        completed = subprocess.run(
            ["git", *args], check=True, capture_output=True, text=True, timeout=3,
        )
        return completed.stdout.strip() or None
    except (OSError, subprocess.SubprocessError):
        return None


def _dependencies() -> dict[str, str | None]:
    values = {}
    for package in _PACKAGES:
        try:
            values[package] = version(package)
        except PackageNotFoundError:
            values[package] = None
    return values


def _flatten(value: Any, prefix: str = "") -> dict[str, Any]:
    if isinstance(value, dict):
        flattened: dict[str, Any] = {}
        for key, child in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            flattened.update(_flatten(child, path))
        return flattened
    if isinstance(value, (list, tuple, set)):
        return {prefix: json.dumps(value, default=str, ensure_ascii=False)}
    return {prefix: value}


def _csv_value(value: Any) -> Any:
    if value is None:
        return "N/A"
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, default=str, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _safe_excel(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, default=str, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    columns = fieldnames or sorted({key for row in rows for key in row})
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _csv_value(row.get(key)) for key in columns})


def _markdown_report(summary: dict[str, Any], results: list[dict[str, Any]]) -> str:
    metadata = summary["metadata"]
    lines = [
        "# Clinical Multi-Agent Benchmark Report",
        "",
        f"- **Executed:** {metadata['executed_at']}",
        f"- **Git:** `{metadata['git']['branch'] or 'N/A'}` / `{metadata['git']['commit'] or 'N/A'}`",
        f"- **Mode:** {metadata['configuration']['mode']}",
        f"- **Dataset:** {metadata['dataset_source']}",
        f"- **Requested documents:** {metadata['configuration']['documents']}",
        f"- **Workers:** {metadata['configuration']['workers']}",
        "",
        "## Environment",
        "",
        f"- **Python:** {metadata['system']['python_version']}",
        f"- **OS:** {metadata['system']['operating_system']}",
        f"- **CPU:** {metadata['system']['cpu'] or 'N/A'} ({metadata['system']['cpu_count'] or 'N/A'} logical processors)",
        f"- **RAM:** {metadata['system']['ram_total_mb'] or 'N/A'} MB",
        f"- **GPU:** {metadata['system']['gpu'] or 'N/A'}",
        "",
        "## Suite Results",
        "",
    ]
    for result in results:
        lines.extend([f"### {result['suite']}", "", "| Metric | Value |", "|---|---:|"])
        for name, value in _flatten(result.get("metrics", {})).items():
            rendered = "N/A" if value is None else str(value)
            escaped = rendered.replace("|", "\\|")
            lines.append(f"| {name} | {escaped} |")
        warnings = result.get("warnings", [])
        if warnings:
            lines.extend(["", "**Warnings / limitations:**"])
            lines.extend(f"- {warning}" for warning in warnings)
        errors = result.get("errors", [])
        if errors:
            lines.extend(["", "**Failed cases:**"])
            lines.extend(f"- `{error.get('document_id', error.get('operation', 'suite'))}`: {error.get('error')}" for error in errors)
        lines.append("")
    lines.extend([
        "## Model Inventory",
        "",
        "Names below are configured in source; their presence does not mean they were loaded during this run.",
        "",
    ])
    lines.extend(f"- **{name}:** {model}" for name, model in metadata["models"].items())
    lines.extend([
        "",
        "## Interpretation",
        "",
        "Accuracy values are only calculated when explicit annotations are present. Engine confidence is not treated as accuracy. This benchmark does not establish clinical safety or eliminate unsupported outputs.",
        "",
    ])
    return "\n".join(lines)


def _write_excel(path: Path, summary: dict[str, Any], results: list[dict[str, Any]],
                 detail_rows: list[dict[str, Any]], stage_rows: list[dict[str, Any]],
                 resource_rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheets = {
        "Summary": [
            {"suite": result["suite"], "metric": key, "value": value}
            for result in results for key, value in _flatten(result.get("metrics", {})).items()
        ],
        "Details": detail_rows,
        "Stage Timings": stage_rows,
        "Resources": resource_rows,
        "Failures": [
            {"suite": result["suite"], **error}
            for result in results for error in result.get("errors", [])
        ],
    }
    for index, (name, rows) in enumerate(sheets.items()):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        columns = sorted({key for row in rows for key in row}) or ["suite", "metric", "value"]
        sheet.append(columns)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24475B")
        for row in rows:
            sheet.append([_safe_excel(row.get(column)) for column in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            width = min(60, max(10, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            sheet.column_dimensions[column_cells[0].column_letter].width = width
    workbook.properties.title = "Clinical Multi-Agent Benchmark"
    workbook.properties.subject = f"Benchmark version {benchmark_version}"
    workbook.save(path)


def generate_reports(
    results: list[BenchmarkResult],
    config: BenchmarkConfig,
    *,
    dataset_source: str = "not configured",
    executed_at: datetime | None = None,
) -> Path:
    """Write all requested report formats and return the unique run directory."""
    timestamp = executed_at or datetime.now().astimezone()
    base = config.output_dir / timestamp.strftime("%Y%m%d_%H%M%S")
    output = base
    suffix = 1
    while output.exists():
        output = base.with_name(f"{base.name}_{suffix:02d}")
        suffix += 1
    output.mkdir(parents=True, exist_ok=False)

    system = system_information()
    summary = {
        "metadata": {
            "executed_at": timestamp.isoformat(),
            "benchmark_framework_version": benchmark_version,
            "git": {"commit": _git_value("rev-parse", "HEAD"), "branch": _git_value("branch", "--show-current")},
            "python_version": system["python_version"],
            "operating_system": system["operating_system"],
            "cpu": system["cpu"],
            "ram_total_mb": system["ram_total_mb"],
            "gpu": system["gpu"],
            "system": system,
            "models": model_inventory(),
            "model_versions": {
                "model_repository_revisions": None,
                "note": "The project configures model identifiers but does not pin or report repository revisions.",
            },
            "dependencies": _dependencies(),
            "dataset_source": dataset_source,
            "configuration": {
                "mode": config.mode, "dataset": str(config.dataset) if config.dataset else None,
                "documents": config.documents, "workers": config.workers,
                "timeout_seconds": config.timeout, "seed": config.seed,
                "output_dir": str(config.output_dir),
            },
            "platform": platform.platform(),
        },
        "suite_summaries": {result.suite: result.metrics for result in results},
        "results": [result.as_dict() for result in results],
        "warnings": [warning for result in results for warning in result.warnings],
        "failed_cases": [
            {"suite": result.suite, **error}
            for result in results for error in result.errors
        ],
    }
    with (output / "summary.json").open("w", encoding="utf-8") as handle:
        json.dump(summary, handle, indent=2, ensure_ascii=False, default=str)

    metric_rows = [
        {"suite": result.suite, "metric": key, "value": value}
        for result in results for key, value in _flatten(result.metrics).items()
    ]
    detail_rows = [
        {"suite": result.suite, **detail}
        for result in results for detail in result.details
    ]
    stage_rows = [row for result in results for row in result.stage_timings]
    resource_rows = [
        {"suite": result.suite, **row}
        for result in results for row in result.resource_metrics
    ]
    _write_csv(output / "summary.csv", metric_rows, ["suite", "metric", "value"])
    _write_csv(output / "detailed_metrics.csv", detail_rows)
    _write_csv(output / "resource_metrics.csv", resource_rows)
    _write_csv(output / "stage_timings.csv", stage_rows)
    (output / "benchmark_report.md").write_text(
        _markdown_report(summary, [result.as_dict() for result in results]),
        encoding="utf-8",
    )
    try:
        _write_excel(output / "benchmark_report.xlsx", summary, [result.as_dict() for result in results],
                     detail_rows, stage_rows, resource_rows)
    except ImportError as exc:
        raise RuntimeError("Excel report generation requires openpyxl; install project requirements.") from exc
    return output
